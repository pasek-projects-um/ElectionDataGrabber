from __future__ import annotations

import csv
import io
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile

from election_data_grabber.adapters.generic_csv import parse_generic_precinct_csv


def _rows_from_xlsx(body: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return [["" if value is None else str(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _rows_from_xls(body: bytes) -> list[list[str]]:
    import xlrd

    with NamedTemporaryFile(suffix=".xls") as tmp:
        tmp.write(body)
        tmp.flush()
        workbook = xlrd.open_workbook(tmp.name)
        sheet = workbook.sheet_by_index(0)
        return [
            ["" if value is None else str(value) for value in sheet.row_values(index)]
            for index in range(sheet.nrows)
        ]


def _looks_like_header(row: list[str]) -> bool:
    tokens = {cell.strip().lower() for cell in row if cell.strip()}
    has_unit = bool(tokens & {"precinct", "reporting_unit", "ward_precinct", "precinct_name", "polling_place"})
    has_contest = bool(tokens & {"office", "contest", "race", "contest_name"})
    has_candidate = bool(tokens & {"candidate", "choice", "candidate_name", "option"})
    return has_unit and has_contest and has_candidate


def _to_csv_bytes(rows: list[list[str]]) -> bytes:
    start = next((i for i, row in enumerate(rows) if _looks_like_header(row)), None)
    if start is None:
        return b""
    out = io.StringIO()
    writer = csv.writer(out)
    for row in rows[start:]:
        writer.writerow(row)
    return out.getvalue().encode("utf-8")


def parse_generic_precinct_excel(
    body: bytes,
    *,
    filename: str,
    election_id: str,
    jurisdiction_id: str,
    source_id: str,
    fetched_at: datetime,
):
    suffix = Path(filename.split("?", 1)[0]).suffix.lower()
    if suffix == ".xlsx":
        rows = _rows_from_xlsx(body)
    elif suffix == ".xls":
        rows = _rows_from_xls(body)
    else:
        raise ValueError(f"unsupported Excel extension: {suffix or '<none>'}")

    csv_body = _to_csv_bytes(rows)
    if not csv_body:
        return []
    return parse_generic_precinct_csv(
        csv_body,
        election_id=election_id,
        jurisdiction_id=jurisdiction_id,
        source_id=source_id,
        fetched_at=fetched_at,
    )
